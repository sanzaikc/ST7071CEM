
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_code_block(doc, code):
    p = doc.add_paragraph(code)
    p.style = 'Normal'
    for run in p.runs:
        run.font.name = 'Courier New'
        run.font.size = Pt(9)

def create_report():
    doc = Document()

    # Title
    title = doc.add_heading('Design, Implementation, and Analysis of a Vertical Search Engine and Document Classification System', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 1. Introduction
    doc.add_heading('1. Introduction', level=1)
    
    doc.add_heading('1.1 Project Context', level=2)
    doc.add_paragraph(
        "In the era of information overload, general-purpose search engines like Google are often too broad for specialized academic research. "
        "Researchers frequently struggle to locate specific outputs from a single department or research group amidst the noise of the global web. "
        "This project addresses this challenge by developing a Vertical Search Engine—a system optimized for a specific domain."
    )
    doc.add_paragraph(
        "The target domain is the Research Centre for Computational Science and Mathematical Modelling at Coventry University. "
        "The goal is to create a centralized, searchable index of all publications (papers, books, articles) where at least one author is a member of this centre. "
        "This system mimics the functionality of Google Scholar but with a highly specific scope, ensuring 100% precision for queries related to this department."
    )

    doc.add_heading('1.2 Scope and Objectives', level=2)
    doc.add_paragraph("The project is divided into two distinct but related computational tasks:")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Information Retrieval (IR)").bold = True
    p.add_run(": Building a full-stack search engine that autonomously crawls, indexes, and retrieves publication data.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Document Classification").bold = True
    p.add_run(": Implementing a supervised machine learning system to automatically classify text into predefined categories (Business, Entertainment, Health) using Multinomial Naive Bayes.")

    # 2. Web Crawler Module
    doc.add_heading('2. Web Crawler Module', level=1)

    doc.add_heading('2.1 Context: The Challenge of Data Acquisition', level=2)
    doc.add_paragraph(
        "Before a search engine can answer queries, it must possess data. Since the target data resides on the web (PurePortal), "
        "manual collection is infeasible due to volume and updates. A Web Crawler (or spider) is required to systematically browse the target website and download content."
    )

    doc.add_heading('2.2 Design Choice: Selenium vs. Requests', level=2)
    p = doc.add_paragraph()
    p.add_run("Why Selenium?").bold = True
    doc.add_paragraph(
        "Early analysis of the PurePortal website revealed that it relies heavily on dynamic content rendering. "
        "Standard HTTP libraries like 'requests' would only retrieve the initial HTML skeleton, missing critical data loaded via JavaScript. "
        "I chose Selenium because it automates a real web browser (Chrome), executing JavaScript just like a human user would. "
        "This ensures complete data capture, including dynamically loaded publication lists."
    )

    doc.add_heading('2.3 The "Politeness" Policy', level=2)
    p = doc.add_paragraph()
    p.add_run("Why is this necessary?").bold = True
    doc.add_paragraph(
        "Web scraping can be aggressive. A crawler hitting a server hundreds of times per second acts like a Denial-of-Service (DoS) attack, "
        "potentially crashing the university's server or getting my IP banned."
    )
    
    p = doc.add_paragraph()
    p.add_run("Implementation:").bold = True
    doc.add_paragraph("To mitigate this, I implemented a strict Politeness Policy:")
    
    p = doc.add_paragraph(style='List Number')
    p.add_run("Robots.txt Compliance:").bold = True
    p.add_run(" The system first checks https://pureportal.coventry.ac.uk/robots.txt to see which paths are allowed or disallowed.")
    
    p = doc.add_paragraph(style='List Number')
    p.add_run("Rate Limiting:").bold = True
    p.add_run(" I introduced a mandatory delay (e.g., 2 seconds) between requests. This 'sleep' period ensures the server load remains negligible.")

    doc.add_paragraph("Code Reference: Robots.txt & Delay (server/crawler/robots.py)", style='Intense Quote')
    code = (
        "# The system parses robots.txt to respect the site owner's rules\n"
        "def robots_allows(url, user_agent='*', timeout_s=10, ...):\n"
        "    # ... logic to check permission ...\n\n"
        "# Application of delay to prevent server overload\n"
        "async def apply_polite_delay(url: str, robots_parser_cache: Dict):\n"
        "    delay = settings.crawl_delay\n"
        "    # ... logic to sleep for 'delay' seconds ...\n"
        "    await asyncio.sleep(delay)"
    )
    add_code_block(doc, code)

    # 3. Discovery Strategy
    doc.add_heading('3. Discovery Strategy: Finding Authors and Publications', level=1)

    doc.add_heading('3.1 Context: Heuristic Traversal', level=2)
    doc.add_paragraph(
        "A crawler needs a map. It cannot blindly visit every link on the internet. "
        "It needs a heuristic to distinguish between a 'relevant' link (a researcher's profile) and an 'irrelevant' one (a 'Privacy Policy' page)."
    )

    doc.add_heading('3.2 Implementation: Regex-Based URL Filtering', level=2)
    p = doc.add_paragraph()
    p.add_run("Why Regex?").bold = True
    doc.add_paragraph(
        "Regular Expressions provide a powerful, pattern-based way to filter text. I analyzed the PurePortal URL structure and identified consistent patterns:"
    )
    doc.add_paragraph("• Author profiles always contain /persons/", style='List Bullet')
    doc.add_paragraph("• Publication pages always contain /publications/", style='List Bullet')
    doc.add_paragraph(
        "By applying these regex filters, the crawler focuses its resources only on the pages that matter, "
        "ignoring navigation links, footers, and external sites. This significantly increases efficiency."
    )

    doc.add_paragraph("Code Reference: Pattern Matching (server/crawler/pureportal.py)", style='Intense Quote')
    code = (
        "def extract_pure_id(url: str) -> Optional[str]:\n"
        "    # Pattern to identify a Publication page\n"
        "    match = re.search(r'/publications/([a-zA-Z0-9-]+)', url)\n"
        "    if match: return match.group(1)\n\n"
        "    # Pattern to identify an Author/Person page\n"
        "    match = re.search(r'/persons/([a-zA-Z0-9-]+)', url)\n"
        "    if match: return match.group(1)"
    )
    add_code_block(doc, code)

    doc.add_heading('3.3 Automated Scheduling', level=2)
    p = doc.add_paragraph()
    p.add_run("Why Schedule?").bold = True
    doc.add_paragraph("Academic records change. New papers are published, and new staff members join. A static database would quickly become obsolete.")
    
    p = doc.add_paragraph()
    p.add_run("Implementation:").bold = True
    doc.add_paragraph(
        "I integrated APScheduler to run the crawler as a background job. "
        "This ensures the search index is 'self-healing' and always up-to-date without requiring human intervention to press a 'start' button."
    )

    doc.add_paragraph("Code Reference: Task Scheduling (server/crawler/scheduler.py)", style='Intense Quote')
    code = (
        "# Configuration for a weekly automated crawl\n"
        "elif settings.crawl_schedule_interval == 'weekly':\n"
        "    trigger = CronTrigger(day_of_week='mon', hour=hour, minute=minute)"
    )
    add_code_block(doc, code)

    # 4. Document Classification
    doc.add_heading('4. Document Classification', level=1)
    
    doc.add_heading('4.1 Context: Supervised Text Classification', level=2)
    doc.add_paragraph(
        "A search engine must do more than just retrieve documents; it must also help users navigate and categorize them. "
        "We implemented a Supervised Text Classification system using Multinomial Naive Bayes to automatically categorize documents "
        "into predefined classes: Business, Entertainment, and Health."
    )
    doc.add_paragraph(
        "Unlike unsupervised clustering (e.g., K-Means), supervised classification uses labeled training data to learn the relationship "
        "between document features and categories. This provides explicit, reproducible predictions with measurable accuracy."
    )

    doc.add_heading('4.2 Implementation: Multinomial Naive Bayes', level=2)
    p = doc.add_paragraph()
    p.add_run("Why Naive Bayes?").bold = True
    doc.add_paragraph(
        "We chose Multinomial Naive Bayes because it is particularly well-suited for text classification tasks. "
        "It applies Bayes' theorem with a 'naive' assumption of feature independence, calculating P(class|document) based on word frequencies. "
        "Despite its simplicity, it performs remarkably well on text data and provides probability estimates for predictions."
    )

    doc.add_heading('4.3 Data Preparation & TF-IDF', level=2)
    doc.add_paragraph(
        "Before classification, text must be converted into numerical form. We used TF-IDF (Term Frequency-Inverse Document Frequency). "
        "This statistic reflects how important a word is to a document in a collection or corpus."
    )
    p = doc.add_paragraph()
    p.add_run("The Logic:").bold = True
    doc.add_paragraph(
        "If a word appears in every document, it is not useful for classification. TF-IDF lowers the weight of such common words "
        "and boosts unique identifiers, ensuring that the classifier learns meaningful patterns for each category."
    )

    doc.add_paragraph("Code Reference: Classification Logic (server/classification/model.py)", style='Intense Quote')
    code = (
        "class ClassifierService:\n"
        "    def train(self):\n"
        "        # TF-IDF Vectorization\n"
        "        self.vectorizer = TfidfVectorizer(stop_words='english')\n"
        "        X_train = self.vectorizer.fit_transform(texts)\n"
        "        \n"
        "        # Train Multinomial Naive Bayes\n"
        "        self.classifier = MultinomialNB(alpha=0.1)\n"
        "        self.classifier.fit(X_train, labels)\n"
        "    \n"
        "    def predict(self, text):\n"
        "        X = self.vectorizer.transform([text])\n"
        "        label = self.classifier.predict(X)[0]\n"
        "        confidence = max(self.classifier.predict_proba(X)[0])\n"
        "        return label, confidence"
    )
    add_code_block(doc, code)

    # 5. Search Engine Core
    doc.add_heading('5. Search Engine Core: Indexing & Retrieval', level=1)

    doc.add_heading('5.1 Theoretical Foundation: The Vector Space Model (VSM)', level=2)
    doc.add_paragraph(
        "At the heart of any modern search engine lies the Vector Space Model. This is an algebraic model for representing text documents (and any objects, in general) "
        "as vectors of identifiers, such as, for example, index terms. It allows us to apply geometric principles to solve the problem of information retrieval."
    )
    p = doc.add_paragraph()
    p.add_run("Why we integrated VSM:").bold = True
    doc.add_paragraph(
        "A simple database query (like SQL 'LIKE %keyword%') is binary: a document either matches or it doesn't. "
        "It cannot tell you which document is *more* relevant. "
        "By integrating VSM, we transform our search problem into a geometry problem. "
        "We can measure the 'distance' or 'angle' between a user's query and every document in our database. "
        "This allows us to rank results by relevance, placing the most important papers at the top of the list—a critical feature for academic research."
    )

    doc.add_heading('5.2 Implementation: From Text to Vectors', level=2)
    doc.add_paragraph(
        "To implement VSM, we constructed a custom pipeline that converts raw text into mathematical objects. "
        "This involves three distinct stages, each addressing a specific linguistic challenge."
    )

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Tokenization (The 'Atomization' of Text)").bold = True
    p.add_run(": We break streams of text into individual words (tokens). We used NLTK's tokenizer instead of simple string splitting to handle punctuation correctly.")

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Term Frequency (TF) Weighting").bold = True
    p.add_run(
        ": We count how often a word appears. However, raw counts can be misleading (a long document naturally has higher counts). "
        "We applied Logarithmic Scaling (1 + log(tf)) to dampen the effect of high-frequency words. "
        "This ensures that mentioning a keyword 100 times doesn't make a document 100 times more relevant than mentioning it once."
    )

    p = doc.add_paragraph(style='List Bullet')
    p.add_run("L2 Normalization").bold = True
    p.add_run(
        ": We normalize all vectors to a unit length. This effectively cancels out document length. "
        "Without this, a 500-page book would always outrank a 10-page paper simply because it has more words."
    )

    doc.add_paragraph("Code Reference: The Math of Relevance (server/indexing/text.py)", style='Intense Quote')
    code = (
        "def tf_weight(tf: int) -> float:\n"
        "    \"\"\"Compute a log-scaled TF weight.\"\"\"\n"
        "    if tf <= 0:\n"
        "        return 0.0\n"
        "    return 1.0 + log(tf)\n\n"
        "def l2_norm(weights: dict[str, float]) -> float:\n"
        "    \"\"\"Compute L2 norm of a sparse weight vector.\"\"\"\n"
        "    return sqrt(sum(w * w for w in weights.values()))"
    )
    add_code_block(doc, code)

    doc.add_heading('5.3 Ranking Algorithm: Cosine Similarity', level=2)
    doc.add_paragraph(
        "Once we have vectors, we need a way to compare them. We chose Cosine Similarity."
    )
    p = doc.add_paragraph()
    p.add_run("The 'Why':").bold = True
    doc.add_paragraph(
        "Cosine similarity measures the cosine of the angle between two vectors projected in a multi-dimensional space. "
        "In the context of text mining, the vectors are arrays of word counts. "
        "If two documents share many similar words in similar proportions, their vectors will point in the same direction, and the angle between them will be small (cosine near 1). "
        "If they share no words, they are orthogonal (cosine 0)."
    )
    doc.add_paragraph(
        "This is superior to Euclidean distance for text because it focuses on the *composition* of the document rather than its *magnitude* (length). "
        "This ensures that a short abstract matches a query just as well as a full textbook if the content is relevant."
    )

    doc.add_paragraph("Code Reference: Cosine Similarity (server/indexing/similarity.py)", style='Intense Quote')
    code = (
        "def cosine_similarity(dot: float, norm_a: float, norm_b: float) -> float:\n"
        "    \"\"\"Compute cosine similarity given dot product and norms.\"\"\"\n"
        "    if norm_a <= 0.0 or norm_b <= 0.0:\n"
        "        return 0.0\n"
        "    return dot / (norm_a * norm_b)"
    )
    add_code_block(doc, code)

    # 6. User Interface
    doc.add_heading('6. User Interface (UI)', level=1)

    doc.add_heading('6.1 Context: Usability', level=2)
    doc.add_paragraph("A powerful backend is useless if the user cannot interact with it. The requirement was to mimic the familiar interface of Google Scholar.")

    doc.add_heading('6.2 Implementation', level=2)
    doc.add_paragraph("I built a web-based frontend using React. It features:")
    doc.add_paragraph("• Instant Feedback: Loading states inform the user the system is working.", style='List Bullet')
    doc.add_paragraph("• Result Cards: displaying Title, Author, Year, and a snippet.", style='List Bullet')
    doc.add_paragraph("• Direct Access: As requested, users can click links to go directly to the PurePortal source, rather than copying/pasting text.", style='List Bullet')

    doc.add_paragraph("Code Reference: UI Components (ui/src/pages/HomePage.tsx)", style='Intense Quote')
    code = (
        "// The main search interface component\n"
        "<SearchBox onSearch={handleSearch} loading={loading} />\n"
        "{results.map((publication, index) => (\n"
        "  // Rendering each result as a distinct card\n"
        "  <PublicationCard key={publication._id} publication={publication} />\n"
        "))}"
    )
    add_code_block(doc, code)

    # 7. Conclusion
    doc.add_heading('7. Conclusion', level=1)
    doc.add_paragraph(
        "This project successfully delivers a robust, vertical search engine that addresses the specific needs of the Coventry University Research Centre. "
        "By combining a polite, scheduled crawler with a TF-IDF/Cosine Similarity ranking engine, the system ensures high-precision retrieval of academic outputs. "
        "Furthermore, the integration of Multinomial Naive Bayes classification demonstrates the system's capability to accurately categorize documents into predefined classes using supervised learning. "
        "The result is a tool that not only retrieves data but classifies it, turning raw information into actionable knowledge."
    )

    # 8. References
    doc.add_heading('8. References', level=1)
    
    # Function to add hanging indent paragraph
    def add_reference(text):
        p = doc.add_paragraph(text)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)

    add_reference("Coventry University. (n.d.). Research Centre for Computational Science and Mathematical Modelling. PurePortal. https://pureportal.coventry.ac.uk/en/organisations/ics-researchcentre-for-computational-science-and-mathematical-mo")
    add_reference("Manning, C. D., Raghavan, P., & Schütze, H. (2008). Introduction to Information Retrieval. Cambridge University Press.")
    add_reference("Bird, S., Klein, E., & Loper, E. (2009). Natural Language Processing with Python. O'Reilly Media.")
    add_reference("Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., Blondel, M., Prettenhofer, P., Weiss, R., Dubourg, V., Vanderplas, J., Passos, A., Cournapeau, D., Brucher, M., Perrot, M., & Duchesnay, E. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825–2830.")
    add_reference("Selenium Project. (2024). The Selenium Browser Automation Project. https://www.selenium.dev/")
    add_reference("Richardson, L. (2007). Beautiful Soup Documentation. https://www.crummy.com/software/BeautifulSoup/bs4/doc/")

    doc.save('Project_Report.docx')

if __name__ == "__main__":
    create_report()
